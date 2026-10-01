#!/usr/bin/env python3
"""Copilot Credits report: Date x Agent x User x Billing Policy x Credits x Cost.

Copilot Studio - Power Platform API (api-version 2024-10-01), read-only:
  GET /licensing/entitlements/MCSMessages/users                     users with consumption on a day
  GET /licensing/entitlements/MCSMessages/users/{userId}/resources  agent x feature credits for that user and day
  GET /licensing/billingPolicies                                    environment -> billing policy, Copilot Studio PAYG flag
  GET /licensing/environments/{environmentId}/entitlements          environment allocation and snapshot totals
  GET /licensing/entitlements/MCSMessages                           tenant capacity totals
  GET /environmentmanagement/environments                           environment names
Microsoft 365 Copilot Chat metered agents - CSV export of the Microsoft 365 admin center Copilot Credits report
('Users and agents' table), passed with --m365-credits-csv. Microsoft offers no API for it.
Enrichment and money - Microsoft Graph $batch (name, department, cost center) and the Azure Cost Management
Query API (actual charges per billing policy per day).

Examples:
  python copilot_credits_report.py --tenant <tenant-guid> --from 2026-09-01 --to 2026-09-30
  python copilot_credits_report.py --tenant <tenant-guid> --m365-credits-csv users_and_agents.csv --m365-as-of 2026-09-30
  python copilot_credits_report.py --no-studio --no-azure --no-graph --m365-as-of 2026-09-30 \\
      --m365-credits-csv samples/m365_credits_users_and_agents_sample.csv
"""
import argparse
import concurrent.futures as cf
import csv
import datetime as dt
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
from collections import defaultdict
from urllib.parse import quote

import msal
import requests
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

__version__ = "1.0.0"
PP, GRAPH, ARM = "https://api.powerplatform.com", "https://graph.microsoft.com", "https://management.azure.com"
ENT, V = "MCSMessages", "2024-10-01"
AZ_CLI_APP = "04b07795-8ddb-461a-bbee-02f9e1bf7b46"
STUDIO, CHAT_ENV = "Copilot Studio", "Microsoft 365 Copilot Chat"
M365 = "Microsoft 365 Copilot Chat (metered agents)"
# Matches current and earlier names of the Copilot Studio meter on Microsoft.PowerPlatform/accounts resources
COPILOT_METER = re.compile(r"copilot|virtual agent|message|credit", re.I)
RATE_SOURCE = ("Source: Microsoft Learn, 'Pay-as-you-go meters', Copilot Studio meter = $0.01 per Copilot Credit "
               "(list price), https://learn.microsoft.com/power-platform/admin/pay-as-you-go-meters. "
               "Change this if your agreement or pre-purchase plan gives a different effective rate.")


class Auth:
    def __init__(self, tenant, mode, client_id, cache_path):
        self.tenant, self.mode, self.lock, self.tokens = tenant, mode, threading.Lock(), {}
        if mode == "devicecode":
            self.cache_path, self.cache = cache_path, msal.SerializableTokenCache()
            if os.path.exists(cache_path):
                with open(cache_path) as f:
                    self.cache.deserialize(f.read())
            self.app = msal.PublicClientApplication(
                client_id, authority=f"https://login.microsoftonline.com/{tenant}", token_cache=self.cache)

    def token(self, resource):
        with self.lock:
            tok, exp = self.tokens.get(resource, (None, 0))
            if not tok or exp - time.time() < 300:
                tok, exp = self._azcli(resource) if self.mode == "azcli" else self._msal(resource)
                self.tokens[resource] = (tok, exp)
            return tok

    def _azcli(self, resource):
        p = subprocess.run([shutil.which("az") or "az", "account", "get-access-token", "--tenant", self.tenant,
                            "--resource", resource, "-o", "json"], capture_output=True, text=True, timeout=180)
        if p.returncode:
            raise SystemExit(f"az account get-access-token failed for {resource}: {p.stderr.strip()[:400]}")
        j = json.loads(p.stdout)
        return j["accessToken"], float(j.get("expires_on") or time.time() + 1800)

    def _msal(self, resource):
        scopes = [resource.rstrip("/") + "/.default"]
        accounts = self.app.get_accounts()
        r = self.app.acquire_token_silent(scopes, account=accounts[0]) if accounts else None
        if not r:
            flow = self.app.initiate_device_flow(scopes=scopes)
            if "user_code" not in flow:
                raise SystemExit(f"Device code flow failed: {flow.get('error_description')}")
            print(flow["message"], file=sys.stderr, flush=True)
            r = self.app.acquire_token_by_device_flow(flow)
        if self.cache.has_state_changed:
            os.makedirs(os.path.dirname(self.cache_path) or ".", exist_ok=True)
            with open(self.cache_path, "w") as f:
                f.write(self.cache.serialize())
        if "access_token" not in r:
            raise SystemExit(f"Token request failed for {resource}: {r.get('error_description')}")
        return r["access_token"], time.time() + int(r.get("expires_in", 3600))


def retry_delay(r, attempt):
    """Seconds to wait before retrying. Azure Cost Management uses x-ms-ratelimit-*-retry-after headers."""
    waits = []
    for k, v in r.headers.items():
        if k.lower().endswith("retry-after"):
            try:
                waits.append(float(v))
            except ValueError:
                pass
    return min(max(waits), 120.0) if waits else float(2 ** attempt)


class Api:
    def __init__(self, auth):
        self.auth, self.s = auth, requests.Session()

    def call(self, method, url, resource, **kw):
        for attempt in range(6):
            r = self.s.request(method, url, timeout=120, **kw,
                               headers={"Authorization": "Bearer " + self.auth.token(resource)})
            if r.status_code not in (429, 500, 502, 503, 504):
                return r
            time.sleep(retry_delay(r, attempt))
        return r

    def pp_paged(self, path, key, **params):
        out, url, params, seen = [], PP + path, dict(params, **{"api-version": V}), set()
        while url:
            r = self.call("GET", url, PP, params=params)
            if r.status_code == 204:
                break
            if r.status_code == 403:
                raise PermissionError(path)
            r.raise_for_status()
            j = r.json()
            for v in j.get("value", []):
                out.extend(v.get(key, []))
            ct, nl = j.get("continuationtoken") or j.get("continuationToken"), j.get("@odata.nextLink")
            if ct and ct not in seen:
                seen.add(ct)
                params["continuationToken"] = ct
            elif nl and nl not in seen:
                seen.add(nl)
                url, params = nl, None
            else:
                url = None
        return out


# ---------------------------------------------------------------- data collection
def daterange(start, end):
    while start <= end:
        yield start
        start += dt.timedelta(days=1)


def collect_usage(api, start, end, workers):
    rows, gaps = [], []
    for day in daterange(start, end):
        d = day.isoformat()
        users = api.pp_paged(f"/licensing/entitlements/{ENT}/users", "users", fromDate=d, toDate=d)
        totals = defaultdict(float)
        for u in users:
            totals[u.get("userId") or ""] += float(u.get("consumed") or 0) + \
                float((u.get("metadata") or {}).get("NonBillableQuantity") or 0)

        def per_user(uid):
            return uid, api.pp_paged(f"/licensing/entitlements/{ENT}/users/{uid}/resources", "resources",
                                     fromDate=d, toDate=d)

        with cf.ThreadPoolExecutor(workers) as ex:
            for uid, res in ex.map(per_user, [u for u in totals if u]):
                got = 0.0
                for x in res:
                    md = x.get("metadata") or {}
                    billed, nonbilled = float(x.get("consumed") or 0), float(md.get("NonBillableQuantity") or 0)
                    got += billed + nonbilled
                    rows.append({"date": (x.get("asOfDate") or d)[:10], "env_id": x.get("environmentId") or "",
                                 "agent_id": x.get("resourceId") or "",
                                 "agent": md.get("ResourceName") or x.get("resourceId") or "(unknown agent)",
                                 "user_id": uid, "feature": md.get("FeatureName") or md.get("Feature") or "",
                                 "product": md.get("ProductName") or md.get("Product") or "",
                                 "billed": billed, "nonbillable": nonbilled})
                if abs(got - totals[uid]) > 0.01:
                    gaps.append((d, uid, totals[uid], got))
        if "" in totals:
            gaps.append((d, "(no user id)", totals[""], 0.0))
        print(f"  {d}: {len([u for u in totals if u])} user(s) with consumption", file=sys.stderr)
    return rows, gaps


def get_policies(api):
    r = api.call("GET", PP + "/licensing/billingPolicies", PP, params={"api-version": V})
    r.raise_for_status()
    pols = []
    for p in r.json().get("value", []):
        envs = p.get("environmentIds")
        if envs is None:
            e = api.call("GET", f"{PP}/licensing/billingPolicies/{p['id']}/environments", PP,
                         params={"api-version": "2022-03-01-preview"})
            envs = [x.get("environmentId") for x in e.json().get("value", [])] if e.ok else []
        bi = p.get("billingInstrument") or {}
        pols.append({"id": p["id"], "name": p["name"], "envs": envs, "subscription": bi.get("subscriptionId"),
                     "resource_id": (bi.get("id") or "").lower(),
                     "mcs_paygo": any(x.get("entitlementId") == ENT and x.get("payAsYouGoState")
                                      for x in p.get("payGoEntitlements", []))})
    return pols


def get_env_names(api):
    r = api.call("GET", PP + "/environmentmanagement/environments", PP, params={"api-version": V})
    return {e["id"]: e.get("displayName") or e["id"] for e in r.json().get("value", [])} if r.ok else {}


def get_env_entitlement(api, env_id):
    r = api.call("GET", f"{PP}/licensing/environments/{env_id}/entitlements", PP, params={"api-version": V})
    if not r.ok:
        return None
    j = r.json()
    for x in (j if isinstance(j, list) else j.get("value", [])):
        if x.get("entitlementId") == ENT:
            e = x.get("entitlement") or {}
            cap, pg = e.get("capacity") or {}, e.get("payGo") or {}
            used = cap.get("consumed") or {}
            return {"allocated": (cap.get("allocated") or {}).get("value") or 0, "prepaid": used.get("value") or 0,
                    "paygo": (pg.get("consumed") or {}).get("value") or 0, "as_of": (used.get("lastUpdatedOn") or "")[:10]}
    return None


def get_tenant_entitlement(api):
    r = api.call("GET", f"{PP}/licensing/entitlements/{ENT}", PP, params={"api-version": V})
    if not r.ok:
        return None
    cap = ((r.json().get("entitlement") or {}).get("capacity")) or {}
    used = cap.get("consumed") or {}
    return {"entitled": (cap.get("entitled") or {}).get("value") or 0,
            "allocated": (cap.get("allocated") or {}).get("value") or 0,
            "consumed": used.get("value") or 0, "type": used.get("consumptionType") or "",
            "as_of": (used.get("lastUpdatedOn") or "")[:10]}


def resolve_users(api, keys):
    """User ID or UPN -> UPN, display name, department and cost center, 20 lookups per Graph $batch call."""
    out, keys = {}, sorted({k for k in keys if k})
    select = "id,userPrincipalName,displayName,department,employeeOrgData"
    for i in range(0, len(keys), 20):
        pending = dict(enumerate(keys[i:i + 20]))
        for attempt in range(4):
            body = {"requests": [{"id": str(n), "method": "GET", "url": f"/users/{quote(k, safe='@')}?$select={select}"}
                                 for n, k in pending.items()]}
            r = api.call("POST", GRAPH + "/v1.0/$batch", GRAPH, json=body)
            if not r.ok:
                print(f"  Graph lookup failed ({r.status_code}); user IDs are shown instead.", file=sys.stderr)
                break
            retry = {}
            for x in r.json().get("responses", []):
                n = int(x["id"])
                if x.get("status") == 200:
                    u = x.get("body") or {}
                    out[pending[n]] = {"upn": u.get("userPrincipalName") or pending[n],
                                       "display": u.get("displayName") or "", "dept": u.get("department") or "",
                                       "cost_center": (u.get("employeeOrgData") or {}).get("costCenter") or ""}
                elif x.get("status") == 429:
                    retry[n] = pending[n]
            if not retry:
                break
            pending = retry
            time.sleep(2 ** attempt)
    return out


def azure_costs(api, pols, extra_subs, start, end):
    names = {p["resource_id"]: p["name"] for p in pols if p["resource_id"]}
    subs = sorted({p["subscription"] for p in pols if p["subscription"]} | set(extra_subs))
    body = {"type": "ActualCost", "timeframe": "Custom",
            "timePeriod": {"from": f"{start}T00:00:00Z", "to": f"{end}T23:59:59Z"},
            "dataset": {"granularity": "Daily",
                        "aggregation": {"cost": {"name": "Cost", "function": "Sum"},
                                        "qty": {"name": "UsageQuantity", "function": "Sum"}},
                        "grouping": [{"type": "Dimension", "name": n} for n in
                                     ("ResourceId", "MeterCategory", "MeterSubCategory", "Meter")],
                        "filter": {"dimensions": {"name": "ResourceType", "operator": "In",
                                                  "values": ["microsoft.powerplatform/accounts"]}}}}
    rows = []
    for sub in subs:
        url = f"{ARM}/subscriptions/{sub}/providers/Microsoft.CostManagement/query?api-version=2023-11-01"
        while url:
            r = api.call("POST", url, ARM, json=body)
            if not r.ok:
                print(f"  Cost Management query failed for {sub}: {r.status_code} {r.text[:200]}", file=sys.stderr)
                break
            p = r.json()["properties"]
            cols = [c["name"] for c in p["columns"]]
            for row in p["rows"]:
                x = dict(zip(cols, row))
                meter = " / ".join(dict.fromkeys(filter(None, (x.get("MeterCategory"), x.get("MeterSubCategory"),
                                                               x.get("Meter")))))
                if COPILOT_METER.search(meter):
                    rid = x["ResourceId"]
                    rows.append({"date": dt.datetime.strptime(str(x["UsageDate"]), "%Y%m%d").date(),
                                 "policy": names.get(rid.lower()) or rid.rsplit("/", 1)[-1], "meter": meter,
                                 "qty": x["UsageQuantity"], "cost": x["Cost"]})
            url = p.get("nextLink")
    return sorted(rows, key=lambda x: (x["date"], x["policy"], x["meter"])), subs


# ---------------------------------------------------------------- Microsoft 365 Copilot Credits report export
M365_HEADERS = {"agentid": "agent_id", "agentname": "agent", "username": "user", "userprincipalname": "user",
                "billingpolicyid": "policy_id", "pastsevendays": 7, "past7days": 7, "pastthirtydays": 30,
                "past30days": 30, "lastactivitydateutc": "last", "lastactivitydate": "last",
                "reportrefreshdate": "refresh"}


def _num(v):
    try:
        return float(str(v or 0).replace(",", "").strip() or 0)
    except ValueError:
        return 0.0


def _date(v):
    v = (v or "").strip()[:10]
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(v, fmt).date()
        except ValueError:
            pass
    return None


def load_policy_map(path):
    if not path:
        return {}
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = [r for r in csv.reader(f) if len(r) > 1 and r[0].strip()]
    return {r[0].strip(): r[1].strip() for r in rows if "policy" not in r[0].lower() and r[0].strip().lower() != "id"}


def load_m365_credits(paths, window, as_of, policy_map):
    rows, concealed = [], False
    for path in paths:
        with open(path, newline="", encoding="utf-8-sig") as f:
            rd = csv.DictReader(f)
            cols = {}
            for h in rd.fieldnames or []:
                key = M365_HEADERS.get(re.sub(r"[^a-z0-9]", "", h.lower()))
                if key is not None:
                    cols.setdefault(key, h)
            if not {"agent", "user", window} <= set(cols):
                raise SystemExit(f"{path}: expected the 'Users and agents' export of the Copilot Credits report with "
                                 f"'Agent name', 'Username' and 'Past {'seven' if window == 7 else '30'} days' columns. "
                                 f"Found: {rd.fieldnames}")
            default_end = as_of or dt.date.fromtimestamp(os.path.getmtime(path))
            if not as_of and "refresh" not in cols:
                print(f"  {os.path.basename(path)}: no report date in the file, using its modified date {default_end}. "
                      "Pass --m365-as-of with the export date to be exact.", file=sys.stderr)

            def get(x, k):
                return (x.get(cols[k]) or "").strip() if k in cols else ""

            for x in rd:
                credits = _num(get(x, window))
                if not credits:
                    continue
                end = _date(get(x, "refresh")) or default_end
                user, agent, pid = get(x, "user"), get(x, "agent") or "(unknown agent)", get(x, "policy_id")
                concealed |= bool(re.fullmatch(r"[0-9A-Fa-f]{32}", user))
                rows.append({"start": end - dt.timedelta(days=window - 1), "end": end, "agent": agent,
                             "agent_id": get(x, "agent_id") or agent, "user": user or "(unknown user)",
                             "policy_id": pid, "policy": policy_map.get(pid) or pid or "(unknown policy)",
                             "credits": credits, "window": window, "last": _date(get(x, "last")),
                             "file": os.path.basename(path)})
    if concealed:
        print("  Microsoft 365 usernames are concealed (hashed). To show names, turn off 'Display concealed user, "
              "group, and site names' in Microsoft 365 admin center > Settings > Org settings > Reports.",
              file=sys.stderr)
    return rows


# ---------------------------------------------------------------- shaping
def funding(billed, nonbilled, pol, alloc):
    if not billed:
        return "Non-billable" if nonbilled else ""
    if pol and pol["mcs_paygo"]:
        return "Prepaid, then PAYG" if alloc else "PAYG"
    return "Prepaid capacity"


def build(rows, m365, names, env_pol, env_ent, people):
    main, feats, facts, agents, users, chat = {}, [], [], {}, {}, []

    def track(kind, product, agent, agent_id, env, policy, user, display, dept):
        a = agents.setdefault((kind, agent_id), {"kind": kind, "product": product, "agent": agent, "env": env,
                                                 "agent_id": agent_id, "policies": set(), "users": set()})
        a["policies"].add(policy)
        a["users"].add(user.lower())
        u = users.setdefault(user.lower(), {"user": user, "display": display, "dept": dept, "agents": set()})
        u["agents"].add((kind, agent_id))

    for r in sorted(rows, key=lambda x: (x["date"], x["agent"].lower(), x["user_id"], x["feature"])):
        env, pol = names.get(r["env_id"], r["env_id"]), env_pol.get(r["env_id"])
        alloc = (env_ent.get(r["env_id"]) or {}).get("allocated") or 0
        who = people.get(r["user_id"]) or {}
        user = who.get("upn") or r["user_id"]
        day = dt.date.fromisoformat(r["date"])
        base = {"date": day, "env": env, "env_id": r["env_id"], "policy": pol["name"] if pol else "(none)",
                "policy_id": pol["id"] if pol else "", "agent": r["agent"], "agent_id": r["agent_id"],
                "user": user, "user_id": r["user_id"], "display": who.get("display", ""),
                "dept": who.get("dept", ""), "cost_center": who.get("cost_center", ""),
                "product": r.get("product") or (CHAT_ENV if env.lower() == CHAT_ENV.lower() else STUDIO)}
        feature = r["feature"] or "(unspecified)"
        feats.append(dict(base, feature=feature, billed=r["billed"], nonbillable=r["nonbillable"]))
        facts.append(dict(base, start=day, end=day, grain="Day", feature=feature, billed=r["billed"],
                          nonbillable=r["nonbillable"], funding=funding(r["billed"], r["nonbillable"], pol, alloc),
                          source="Power Platform licensing API"))
        m = main.setdefault((r["date"], r["env_id"], r["agent_id"], r["user_id"]),
                            dict(base, billed=0.0, nonbillable=0.0, _pol=pol, _alloc=alloc))
        m["billed"] += r["billed"]
        m["nonbillable"] += r["nonbillable"]
        track("studio", base["product"], r["agent"], r["agent_id"], env, base["policy"], user, base["display"],
              base["dept"])
    main_rows = list(main.values())
    for m in main_rows:
        m["funding"] = funding(m["billed"], m["nonbillable"], m.pop("_pol"), m.pop("_alloc"))

    for x in sorted(m365, key=lambda x: (x["end"], x["agent"].lower(), x["user"].lower())):
        who = people.get(x["user"]) or {}
        c = dict(x, display=who.get("display", ""), dept=who.get("dept", ""), cost_center=who.get("cost_center", ""))
        chat.append(c)
        facts.append({"date": x["end"], "start": x["start"], "end": x["end"], "grain": f"{x['window']} days",
                      "agent": x["agent"], "agent_id": x["agent_id"], "user": x["user"], "user_id": "",
                      "display": c["display"], "dept": c["dept"], "cost_center": c["cost_center"],
                      "policy": x["policy"], "policy_id": x["policy_id"], "product": M365, "feature": "",
                      "billed": x["credits"], "nonbillable": 0.0, "funding": "PAYG (Microsoft 365 billing policy)",
                      "source": "Microsoft 365 admin center - Copilot Credits report", "env": "", "env_id": ""})
        track("m365", M365, x["agent"], x["agent_id"], "", x["policy"], x["user"], c["display"], c["dept"])

    return {"main": main_rows, "features": feats, "facts": facts, "agents": agents, "users": users, "m365": chat,
            "products": sorted({m["product"] for m in main_rows}),
            "fundings": sorted({m["funding"] for m in main_rows} - {"", "Non-billable"})}


def allocate(facts, azure):
    """Split each billing policy's actual Azure cost for a day (or report window) across rows by billed credits."""
    by_policy = defaultdict(list)
    for x in azure:
        by_policy[x["policy"].lower()].append((x["date"], x["cost"]))
    key = lambda f: (f["policy"].lower(), f["start"], f["end"])  # noqa: E731
    credit_totals, pools = defaultdict(float), {}
    for f in facts:
        credit_totals[key(f)] += f["billed"]
    for f in facts:
        k = key(f)
        if k not in pools:
            pools[k] = sum(c for d, c in by_policy.get(k[0], []) if k[1] <= d <= k[2])
        f["allocated"] = pools[k] * f["billed"] / credit_totals[k] if f["billed"] and credit_totals[k] else 0.0


# ---------------------------------------------------------------- Excel output
NUM, INT, USD, DATE = '#,##0.00;(#,##0.00);"-"', '#,##0;(#,##0);"-"', '"$"#,##0.00;("$"#,##0.00);"-"', "yyyy-mm-dd"
BLUE, GREEN, HDR = "0000FF", "008000", PatternFill("solid", start_color="1F4E78")
MAIN, CHAT, FEAT, RECON, RATE = "Agent-User Daily", "M365 Copilot Chat", "Feature Detail", "Reconciliation", \
    "Settings!$B$3"
MC = {k: get_column_letter(i) for i, k in enumerate(("date", "agent", "user", "policy", "credits", "cost", "alloc",
      "billed", "nonb", "env", "funding", "dept", "product", "agent_id", "user_id", "env_id"), 1)}
CC = {k: get_column_letter(i) for i, k in enumerate(("start", "end", "agent", "user", "policy", "credits", "cost",
      "alloc", "window", "last", "dept", "agent_id", "policy_id", "file"), 1)}
FC = {k: get_column_letter(i) for i, k in enumerate(("date", "env", "agent", "user", "feature", "credits", "billed",
      "nonb", "cost", "agent_id", "user_id"), 1)}
SETTINGS_NOTES = [
    "Estimated Cost = Billed Credits x rate. Non-billable credits (for example Microsoft 365 Copilot licensed users or "
    "developer environments) cost nothing.",
    "Allocated Azure Cost = the billing policy's actual Azure charge for the day (or for the Microsoft 365 report "
    "window), split across report rows in proportion to their billed credits.",
    "Funding Source: PAYG = billed to the Azure subscription of the environment's billing policy. Prepaid, then PAYG = "
    "the environment has prepaid credits allocated and overflows to PAYG. Prepaid capacity = covered by capacity "
    "packs, no Azure charge. Non-billable = zero-rated usage.",
    "Microsoft 365 Copilot Chat rows come from the Copilot Credits report export and cover a 7- or 30-day window "
    "ending on 'Period End', not a single day.",
]
ABOUT = [
    "How this report is built",
    "Copilot Studio: Power Platform API (api-version 2024-10-01) - GET /licensing/entitlements/MCSMessages/users and "
    "GET /licensing/entitlements/MCSMessages/users/{userId}/resources, queried one day at a time (fromDate = toDate). "
    "'Billed Credits' is the API 'consumed' value; 'Non-billable Credits' is metadata.NonBillableQuantity.",
    "Billing policy and Copilot Studio PAYG flag: GET /licensing/billingPolicies (payGoEntitlements[MCSMessages]."
    "payAsYouGoState). Environment names: GET /environmentmanagement/environments. Users: Microsoft Graph $batch.",
    "Microsoft 365 Copilot Chat metered agents: the 'Users and agents' CSV export of the Microsoft 365 admin center "
    "Copilot Credits report (Reports > Usage > Microsoft Copilot > Credits). Microsoft offers no API for it.",
    "Actual charges: Azure Cost Management Query API, resource type Microsoft.PowerPlatform/accounts (one resource per "
    "billing policy), Copilot meters only. Azure has no agent or user dimension, so it is allocated to report rows.",
    "Limits: Microsoft restricted the agent listing endpoint (/licensing/entitlements/MCSMessages/resources) to the "
    "Power Platform admin center in September 2026. Usage with no user identity can't be listed per agent and shows up "
    "as Azure cost that is not matched to a report row (see Overview).",
    "The user-level endpoints are in the public REST reference, but Microsoft has said this consumption API family is "
    "intended for the Power Platform admin center. Keep the admin center downloads (Licensing > Copilot Studio > "
    "Download report) as the fallback.",
    "Latency: licensing snapshots are daily (00:00 UTC) and Azure charges can take 24 hours. Report on days that are "
    "at least 2 days old. Validate totals against your Azure invoice before using them for chargeback.",
]


def put(ws, r, c, v, fmt=None, color="000000", bold=False, wrap=False):
    cell = ws.cell(r, c, v)
    cell.font = Font(name="Arial", size=10, color=color, bold=bold)
    if fmt:
        cell.number_format = fmt
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return cell


def header(ws, row, spec):
    for c, (name, width) in enumerate(spec, 1):
        cell = put(ws, row, c, name, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if width:
            ws.column_dimensions[get_column_letter(c)].width = width


def rng(sheet, col, last):
    return f"'{sheet}'!${col}$2:${col}${last}"


def is_link(v):
    return isinstance(v, str) and v.startswith("=") and "!" in v


def totals(ws, first, row, fmts):
    put(ws, row, 1, "Total", bold=True)
    for c, fmt in fmts.items():
        col = get_column_letter(c)
        put(ws, row, c, f"=SUM({col}{first}:{col}{row - 1})" if row > first else 0, fmt, bold=True)


def write_workbook(path, a):
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    main, chat, feats, n_az = a["main"], a["m365"], a["features"], len(a["azure"])
    last, clast, flast = max(2, len(main) + 1), max(2, len(chat) + 1), max(2, len(feats) + 1)

    def M(k):
        return rng(MAIN, MC[k], last)

    def C(k):
        return rng(CHAT, CC[k], clast)

    def F(k):
        return rng(FEAT, FC[k], flast)

    def az(col):
        return f"'{RECON}'!${col}$3:${col}${n_az + 2}"

    ws = wb.active
    ws.title = "Settings"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 34, 60
    put(ws, 1, 1, "Copilot Credits report", bold=True)
    for r, (label, value, fmt, color) in enumerate([
            ("Rate per Copilot Credit (USD)", a["rate"], '"$"#,##0.0000', BLUE),
            ("Report period - from", a["start"], DATE, "000000"),
            ("Report period - to", a["end"], DATE, "000000"),
            ("Tenant ID", a["tenant"] or "(offline)", None, "000000"),
            ("Generated (UTC)", a["generated"], None, "000000"),
            ("Script version", __version__, None, "000000")], 3):
        put(ws, r, 1, label)
        put(ws, r, 2, value, fmt, color)
    ws["B3"].comment = Comment(RATE_SOURCE, "copilot-credits-report")
    for r, text in enumerate(SETTINGS_NOTES, 10):
        put(ws, r, 1, text, wrap=True)
        ws.merge_cells(f"A{r}:B{r}")
        ws.row_dimensions[r].height = 42

    ws = wb.create_sheet("Overview")
    for col, w in zip("ABCDEF", (48, 14, 14, 14, 14, 16)):
        ws.column_dimensions[col].width = w
    put(ws, 1, 1, "Credits and cost by product", bold=True)
    header(ws, 2, [(h, None) for h in ("Product", "Billed Credits", "Non-billable Credits", "Credits Consumed",
                                       "Estimated Cost (USD)", "Allocated Azure Cost (USD)")])
    r = first = 3
    for p in a["products"]:
        k = f"{M('product')},$A{r}"
        for c, v, fmt in ((1, p, None), (2, f"=SUMIFS({M('billed')},{k})", NUM), (3, f"=SUMIFS({M('nonb')},{k})", NUM),
                          (4, f"=B{r}+C{r}", NUM), (5, f"=SUMIFS({M('cost')},{k})", USD),
                          (6, f"=SUMIFS({M('alloc')},{k})", USD)):
            put(ws, r, c, v, fmt, GREEN if is_link(v) else "000000")
        r += 1
    if chat:
        for c, v, fmt in ((1, M365, None), (2, f"=SUM({C('credits')})", NUM), (3, 0, NUM), (4, f"=B{r}+C{r}", NUM),
                          (5, f"=SUM({C('cost')})", USD), (6, f"=SUM({C('alloc')})", USD)):
            put(ws, r, c, v, fmt, GREEN if is_link(v) else "000000")
        r += 1
    if r == first:
        put(ws, r, 1, "No consumption in the report period.")
        r += 1
    totals(ws, first, r, {2: NUM, 3: NUM, 4: NUM, 5: USD, 6: USD})
    product_total = r

    r += 2
    put(ws, r, 1, "Billed credits by funding source", bold=True)
    header(ws, r + 1, [(h, None) for h in ("Funding Source", "Billed Credits", "Estimated Cost (USD)",
                                           "Allocated Azure Cost (USD)")])
    r = first = r + 2
    for fs in a["fundings"]:
        k = f"{M('funding')},$A{r}"
        for c, v, fmt in ((1, fs, None), (2, f"=SUMIFS({M('billed')},{k})", NUM), (3, f"=SUMIFS({M('cost')},{k})", USD),
                          (4, f"=SUMIFS({M('alloc')},{k})", USD)):
            put(ws, r, c, v, fmt, GREEN if is_link(v) else "000000")
        r += 1
    if chat:
        for c, v, fmt in ((1, "PAYG (Microsoft 365 billing policy)", None), (2, f"=SUM({C('credits')})", NUM),
                          (3, f"=SUM({C('cost')})", USD), (4, f"=SUM({C('alloc')})", USD)):
            put(ws, r, c, v, fmt, GREEN if is_link(v) else "000000")
        r += 1
    if r == first:
        put(ws, r, 1, "No billed credits in the report period.")
        r += 1
    totals(ws, first, r, {2: NUM, 3: USD, 4: USD})

    r += 2
    put(ws, r, 1, "Azure actual charges on Copilot meters", bold=True)
    for label, v, fmt in (("Azure cost (USD)", f"=SUM({az('E')})" if n_az else 0, USD),
                          ("Azure quantity (credits)", f"=SUM({az('D')})" if n_az else 0, NUM),
                          ("Allocated to report rows (USD)", f"=F{product_total}", USD)):
        r += 1
        put(ws, r, 1, label)
        put(ws, r, 2, v, fmt, GREEN if is_link(v) else "000000")
    r += 1
    put(ws, r, 1, "Not matched to any report row (USD)")
    put(ws, r, 2, f"=B{r - 3}-B{r - 1}", USD)
    r += 1
    put(ws, r, 1, "Unmatched cost usually means usage with no user (for example autonomous runs), a Microsoft 365 "
                  "billing policy without --m365-policy-map, or metered usage this report doesn't cover.", wrap=True)
    ws.merge_cells(f"A{r}:F{r}")
    ws.row_dimensions[r].height = 30

    ws = wb.create_sheet(MAIN)
    header(ws, 1, [("Date", 12), ("Agent Name", 30), ("User", 36), ("Billing Policy", 22), ("Credits Consumed", 13),
                   ("Estimated Cost (USD)", 13), ("Allocated Azure Cost (USD)", 14), ("Billed Credits", 12),
                   ("Non-billable Credits", 13), ("Environment", 20), ("Funding Source", 18), ("Department", 18),
                   ("Product", 22), ("Agent ID", 38), ("User ID", 38), ("Environment ID", 38)])
    for i, x in enumerate(main, 2):
        alloc = (f"=IFERROR(SUMIFS({az('E')},{az('A')},$A{i},{az('B')},$D{i})*$H{i}"
                 f"/SUMIFS({M('billed')},{M('date')},$A{i},{M('policy')},$D{i}),0)") if n_az else 0
        for c, v in enumerate([x["date"], x["agent"], x["user"], x["policy"], f"=H{i}+I{i}", f"=H{i}*{RATE}", alloc,
                               x["billed"], x["nonbillable"], x["env"], x["funding"], x["dept"], x["product"],
                               x["agent_id"], x["user_id"], x["env_id"]], 1):
            put(ws, i, c, v, {1: DATE, 5: NUM, 6: USD, 7: USD, 8: NUM, 9: NUM}.get(c), GREEN if is_link(v) else "000000")
    if not main:
        put(ws, 2, 1, "No Copilot Studio consumption was returned for the report period.")
    ws.freeze_panes, ws.auto_filter.ref = "A2", f"A1:P{last}"

    if chat:
        ws = wb.create_sheet(CHAT)
        header(ws, 1, [("Period Start", 12), ("Period End", 12), ("Agent Name", 30), ("User", 36),
                       ("Billing Policy", 24), ("Credits", 12), ("Estimated Cost (USD)", 13),
                       ("Allocated Azure Cost (USD)", 14), ("Window (days)", 9), ("Last Activity (UTC)", 13),
                       ("Department", 18), ("Agent ID", 38), ("Billing Policy ID", 38), ("Source File", 30)])
        for i, x in enumerate(chat, 2):
            alloc = (f"=IFERROR(SUMIFS({az('E')},{az('B')},$E{i},{az('A')},\">=\"&$A{i},{az('A')},\"<=\"&$B{i})*$F{i}"
                     f"/SUMIFS({C('credits')},{C('policy')},$E{i},{C('start')},$A{i},{C('end')},$B{i}),0)") \
                if n_az else 0
            for c, v in enumerate([x["start"], x["end"], x["agent"], x["user"], x["policy"], x["credits"],
                                   f"=F{i}*{RATE}", alloc, x["window"], x["last"], x["dept"], x["agent_id"],
                                   x["policy_id"], x["file"]], 1):
                put(ws, i, c, v, {1: DATE, 2: DATE, 6: NUM, 7: USD, 8: USD, 9: INT, 10: DATE}.get(c),
                    GREEN if is_link(v) else "000000")
        ws.freeze_panes, ws.auto_filter.ref = "A2", f"A1:N{clast}"

    ws = wb.create_sheet(FEAT)
    header(ws, 1, [("Date", 12), ("Environment", 20), ("Agent Name", 30), ("User", 36), ("Feature", 24),
                   ("Credits Consumed", 13), ("Billed Credits", 12), ("Non-billable Credits", 13),
                   ("Estimated Cost (USD)", 13), ("Agent ID", 38), ("User ID", 38)])
    for i, x in enumerate(feats, 2):
        for c, v in enumerate([x["date"], x["env"], x["agent"], x["user"], x["feature"], f"=G{i}+H{i}", x["billed"],
                               x["nonbillable"], f"=G{i}*{RATE}", x["agent_id"], x["user_id"]], 1):
            put(ws, i, c, v, {1: DATE, 6: NUM, 7: NUM, 8: NUM, 9: USD}.get(c), GREEN if is_link(v) else "000000")
    ws.freeze_panes, ws.auto_filter.ref = "A2", f"A1:K{flast}"

    ws = wb.create_sheet("Agent Summary")
    header(ws, 1, [("Product", 26), ("Agent Name", 30), ("Environment", 20), ("Billing Policy", 24),
                   ("Credits Consumed", 13), ("Billed Credits", 12), ("Non-billable Credits", 13),
                   ("Estimated Cost (USD)", 13), ("Allocated Azure Cost (USD)", 14), ("Distinct Users", 10),
                   ("Agent ID", 38)])
    agents = sorted(a["agents"].values(), key=lambda x: (x["product"], x["agent"].lower()))
    for i, g in enumerate(agents, 2):
        pol = next(iter(g["policies"])) if len(g["policies"]) == 1 else "(multiple)"
        for c, v in ((1, g["product"]), (2, g["agent"]), (3, g["env"]), (4, pol), (11, g["agent_id"])):
            put(ws, i, c, v)
        put(ws, i, 10, len(g["users"]), INT)
        if g["kind"] == "studio":
            k = f"{M('agent_id')},$K{i}"
            f = {5: f"=SUMIFS({M('credits')},{k})", 6: f"=SUMIFS({M('billed')},{k})", 7: f"=SUMIFS({M('nonb')},{k})",
                 8: f"=SUMIFS({M('cost')},{k})", 9: f"=SUMIFS({M('alloc')},{k})"}
        else:
            k = f"{C('agent_id')},$K{i}"
            f = {5: f"=SUMIFS({C('credits')},{k})", 6: f"=E{i}", 7: 0, 8: f"=SUMIFS({C('cost')},{k})",
                 9: f"=SUMIFS({C('alloc')},{k})"}
        for c, v in f.items():
            put(ws, i, c, v, USD if c in (8, 9) else NUM, GREEN if is_link(v) else "000000")
    totals(ws, 2, len(agents) + 2, {5: NUM, 6: NUM, 7: NUM, 8: USD, 9: USD})
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("User Summary")
    header(ws, 1, [("User", 36), ("Display Name", 24), ("Department", 18), ("Copilot Studio Credits", 13),
                   ("M365 Copilot Chat Credits", 13), ("Billed Credits", 12), ("Non-billable Credits", 13),
                   ("Estimated Cost (USD)", 13), ("Allocated Azure Cost (USD)", 14), ("Agents Used", 10)])
    people = sorted(a["users"].values(), key=lambda x: x["user"].lower())
    for i, u in enumerate(people, 2):
        for c, v in ((1, u["user"]), (2, u["display"]), (3, u["dept"])):
            put(ws, i, c, v)
        put(ws, i, 10, len(u["agents"]), INT)
        s, ck = f"{M('user')},$A{i}", f"{C('user')},$A{i}"

        def plus(col):
            return f"+SUMIFS({C(col)},{ck})" if chat else ""

        f = {4: f"=SUMIFS({M('credits')},{s})", 5: f"=SUMIFS({C('credits')},{ck})" if chat else 0,
             6: f"=SUMIFS({M('billed')},{s})+E{i}", 7: f"=SUMIFS({M('nonb')},{s})",
             8: f"=SUMIFS({M('cost')},{s}){plus('cost')}", 9: f"=SUMIFS({M('alloc')},{s}){plus('alloc')}"}
        for c, v in f.items():
            put(ws, i, c, v, USD if c in (8, 9) else NUM, GREEN if is_link(v) else "000000")
    totals(ws, 2, len(people) + 2, {4: NUM, 5: NUM, 6: NUM, 7: NUM, 8: USD, 9: USD})
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Feature Summary")
    header(ws, 1, [("Feature", 28), ("Credits Consumed", 13), ("Billed Credits", 12), ("Non-billable Credits", 13),
                   ("Estimated Cost (USD)", 13)])
    feature_names = sorted({x["feature"] for x in feats})
    for i, name in enumerate(feature_names, 2):
        put(ws, i, 1, name)
        k = f"{F('feature')},$A{i}"
        for c, col, fmt in ((2, "credits", NUM), (3, "billed", NUM), (4, "nonb", NUM), (5, "cost", USD)):
            put(ws, i, c, f"=SUMIFS({F(col)},{k})", fmt, GREEN)
    totals(ws, 2, len(feature_names) + 2, {2: NUM, 3: NUM, 4: NUM, 5: USD})
    ws.freeze_panes = "A2"

    ws = wb.create_sheet(RECON)
    put(ws, 1, 1, "A. Azure Cost Management - actual charges on Copilot meters, per billing policy per day", bold=True)
    header(ws, 2, [("Date", 22), ("Billing Policy (Azure resource)", 30), ("Meter", 40), ("Azure Quantity (credits)", 14),
                   ("Azure Cost (USD)", 14), ("Report Billed Credits - Copilot Studio (same policy & date)", 22),
                   ("Difference (credits)", 14)])
    r = 3
    for x in a["azure"]:
        for c, v, fmt in ((1, x["date"], DATE), (2, x["policy"], None), (3, x["meter"], None), (4, x["qty"], NUM),
                          (5, x["cost"], USD)):
            put(ws, r, c, v, fmt)
        put(ws, r, 6, f"=SUMIFS({M('billed')},{M('date')},$A{r},{M('policy')},$B{r})", NUM, GREEN)
        put(ws, r, 7, f"=D{r}-F{r}", NUM)
        r += 1
    if not n_az:
        put(ws, r, 1, a["azure_note"])
        r += 1
    if chat:
        r += 1
        put(ws, r, 1, "A2. Microsoft 365 billing policies - Copilot Credits report window vs Azure", bold=True)
        header(ws, r + 1, [(h, None) for h in ("Billing Policy", "Period Start", "Period End",
                                               "Report Credits (M365 export)", "Azure Quantity (same policy & window)",
                                               "Difference (credits)")])
        r += 2
        e = n_az + 2
        for pol, s, end in sorted({(x["policy"], x["start"], x["end"]) for x in chat}, key=lambda t: (t[2], t[0])):
            put(ws, r, 1, pol)
            put(ws, r, 2, s, DATE)
            put(ws, r, 3, end, DATE)
            put(ws, r, 4, f"=SUMIFS({C('credits')},{C('policy')},$A{r},{C('start')},$B{r},{C('end')},$C{r})", NUM,
                GREEN)
            put(ws, r, 5, (f"=SUMIFS($D$3:$D${e},$B$3:$B${e},$A{r},$A$3:$A${e},\">=\"&$B{r},"
                           f"$A$3:$A${e},\"<=\"&$C{r})") if n_az else 0, NUM)
            put(ws, r, 6, f"=E{r}-D{r}", NUM)
            r += 1
    if a["envs"]:
        r += 1
        put(ws, r, 1, "B. Environments (Power Platform licensing API)", bold=True)
        header(ws, r + 1, [(h, None) for h in ("Environment", "Billing Policy", "Copilot Studio PAYG enabled on policy",
                                               "Prepaid Credits Allocated", "Billed Credits - prepaid (last snapshot)",
                                               "Billed Credits - PAYG (last snapshot)", "Snapshot Date")])
        r += 2
        for env in a["envs"]:
            for c, v in enumerate([env["env"], env["policy"], env["paygo"], env["allocated"], env["prepaid"],
                                   env["paygo_used"], env["as_of"]], 1):
                put(ws, r, c, v, NUM if c in (4, 5, 6) else None)
            r += 1
    t = a["tenant_ent"]
    if t:
        r += 1
        put(ws, r, 1, "C. Tenant Copilot Credit capacity (Power Platform licensing API)", bold=True)
        for label, v, fmt in (("Prepaid credits entitled", t["entitled"], NUM),
                              ("Credits allocated to environments", t["allocated"], NUM),
                              (f"Billed credits consumed ({t['type']})", t["consumed"], NUM), ("As of", t["as_of"], None)):
            r += 1
            put(ws, r, 1, label)
            put(ws, r, 2, v, fmt)
    if a["gaps"]:
        r += 2
        put(ws, r, 1, "D. Completeness check - user totals that did not match the agent/feature rows", bold=True)
        header(ws, r + 1, [(h, None) for h in ("Date", "User ID", "User total (credits)", "Sum of agent rows")])
        r += 1
        for g in a["gaps"]:
            r += 1
            for c, v in enumerate(g, 1):
                put(ws, r, c, v, NUM if c > 2 else None)

    ws = wb.create_sheet("About")
    ws.column_dimensions["A"].width = 150
    for i, line in enumerate(ABOUT, 1):
        put(ws, i, 1, line, bold=i == 1, wrap=True)
    wb.save(path)


FACT_COLS = ["Date", "Agent Name", "User", "Billing Policy", "Credits Consumed", "Estimated Cost (USD)",
             "Allocated Azure Cost (USD)", "Product", "Feature", "Billed Credits", "Non-billable Credits",
             "Funding Source", "Grain", "Period Start", "Period End", "Environment", "Department", "Cost Center",
             "Source", "Agent ID", "User ID", "Environment ID", "Billing Policy ID"]


def write_facts(path, facts, rate):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(FACT_COLS)
        for x in facts:
            w.writerow([x["date"], x["agent"], x["user"], x["policy"], round(x["billed"] + x["nonbillable"], 4),
                        round(x["billed"] * rate, 4), round(x["allocated"], 4), x["product"], x["feature"],
                        x["billed"], x["nonbillable"], x["funding"], x["grain"], x["start"], x["end"], x["env"],
                        x["dept"], x["cost_center"], x["source"], x["agent_id"], x["user_id"], x["env_id"],
                        x["policy_id"]])


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--tenant", help="Microsoft Entra tenant ID (not needed with --no-studio --no-azure --no-graph)")
    ap.add_argument("--from", dest="start", type=dt.date.fromisoformat, help="First day, YYYY-MM-DD (default: --to minus 29 days)")
    ap.add_argument("--to", dest="end", type=dt.date.fromisoformat, help="Last day, YYYY-MM-DD (default: yesterday)")
    ap.add_argument("--rate", type=float, default=0.01, help="USD per Copilot Credit (default: 0.01 list price)")
    ap.add_argument("--out", help="Output .xlsx path. The credits table CSV is written next to it.")
    ap.add_argument("--m365-credits-csv", action="append", default=[], metavar="FILE",
                    help="'Users and agents' export of the Microsoft 365 admin center Copilot Credits report (repeatable)")
    ap.add_argument("--m365-window", type=int, choices=(7, 30), default=7,
                    help="Column to use from that export: 7 = 'Past seven days' (default), 30 = 'Past 30 days'")
    ap.add_argument("--m365-as-of", type=dt.date.fromisoformat, metavar="YYYY-MM-DD",
                    help="Date the Microsoft 365 report was exported, which is the end of its window")
    ap.add_argument("--m365-policy-map", metavar="FILE",
                    help="CSV with 'Billing Policy ID,Billing Policy Name' rows; use the Azure resource names")
    ap.add_argument("--auth", choices=["devicecode", "azcli"], default="devicecode")
    ap.add_argument("--client-id", default=AZ_CLI_APP, help="Public client app ID for device code sign-in")
    ap.add_argument("--cache", default=os.path.join(os.path.expanduser("~"), ".copilot-credits-report", "token_cache.bin"))
    ap.add_argument("--subscription", action="append", default=[],
                    help="Extra Azure subscription to read costs from, e.g. the one holding Microsoft 365 billing policies")
    ap.add_argument("--no-studio", action="store_true", help="Skip Copilot Studio (Power Platform API)")
    ap.add_argument("--no-azure", action="store_true", help="Skip Azure Cost Management")
    ap.add_argument("--no-graph", action="store_true", help="Skip Microsoft Graph user lookups")
    ap.add_argument("--workers", type=int, default=8)
    ap.add_argument("--version", action="version", version=__version__)
    o = ap.parse_args()
    offline = o.no_studio and o.no_azure and o.no_graph
    if not o.tenant and not offline:
        ap.error("--tenant is required unless --no-studio, --no-azure and --no-graph are all set")
    if o.no_studio and not o.m365_credits_csv:
        ap.error("--no-studio needs at least one --m365-credits-csv")
    o.end = o.end or dt.date.today() - dt.timedelta(days=1)
    o.start = o.start or o.end - dt.timedelta(days=29)
    o.out = o.out or f"copilot_credits_{o.start}_{o.end}.xlsx"
    api = None if offline else Api(Auth(o.tenant, o.auth, o.client_id, o.cache))

    pols, names, rows, gaps, env_ent, tenant_ent = [], {}, [], [], {}, None
    if not (o.no_studio and o.no_azure):
        print("Reading billing policies...", file=sys.stderr)
        try:
            pols = get_policies(api)
        except requests.RequestException as e:
            print(f"  Could not read billing policies ({e}); rows show '(none)'.", file=sys.stderr)
    env_pol = {e: p for p in pols for e in p["envs"]}
    if not o.no_studio:
        names = get_env_names(api)
        print(f"Reading Copilot Studio credits {o.start} -> {o.end}...", file=sys.stderr)
        try:
            rows, gaps = collect_usage(api, o.start, o.end, o.workers)
        except PermissionError as e:
            raise SystemExit(f"403 Forbidden on {e}. Microsoft may have restricted this endpoint. Use the Power Platform "
                             "admin center downloads (Licensing > Copilot Studio > Download report) instead.")
        env_ent = {e: get_env_entitlement(api, e) for e in sorted({r["env_id"] for r in rows} | set(env_pol))}
        tenant_ent = get_tenant_entitlement(api)
    m365 = load_m365_credits(o.m365_credits_csv, o.m365_window, o.m365_as_of, load_policy_map(o.m365_policy_map))
    people = {}
    if not o.no_graph:
        print("Resolving users in Microsoft Graph...", file=sys.stderr)
        people = resolve_users(api, {r["user_id"] for r in rows} | {x["user"] for x in m365})

    data = build(rows, m365, names, env_pol, env_ent, people)
    azure, subs = [], []
    if not o.no_azure:
        az_start = min([o.start] + [x["start"] for x in m365])
        az_end = max([o.end] + [x["end"] for x in m365])
        print(f"Reading Azure Cost Management {az_start} -> {az_end}...", file=sys.stderr)
        azure, subs = azure_costs(api, pols, o.subscription, az_start, az_end)
    allocate(data["facts"], azure)
    note = ("Azure Cost Management skipped (--no-azure)." if o.no_azure else
            f"No Copilot meter charges found in Azure Cost Management (subscriptions checked: {', '.join(subs) or 'none'}).")
    envs = [{"env": names.get(e, e), "policy": env_pol[e]["name"] if e in env_pol else "(none)",
             "paygo": "Yes" if e in env_pol and env_pol[e]["mcs_paygo"] else "No",
             **{k: (env_ent.get(e) or {}).get(k, "") for k in ("allocated", "prepaid", "as_of")},
             "paygo_used": (env_ent.get(e) or {}).get("paygo", "")} for e in env_ent]

    write_workbook(o.out, dict(data, rate=o.rate, start=o.start, end=o.end, tenant=o.tenant, azure=azure,
                               azure_note=note, envs=envs, tenant_ent=tenant_ent, gaps=gaps,
                               generated=dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M")))
    csv_path = os.path.splitext(o.out)[0] + ".csv"
    write_facts(csv_path, data["facts"], o.rate)

    billed = sum(f["billed"] for f in data["facts"])
    print(f"\nCopilot Studio rows: {len(data['main'])} | Microsoft 365 rows: {len(data['m365'])} | agents: "
          f"{len(data['agents'])} | users: {len(data['users'])} | billed credits: {billed:,.2f} | est. cost: "
          f"${billed * o.rate:,.2f} | Azure cost: ${sum(x['cost'] for x in azure):,.2f} | completeness gaps: {len(gaps)}"
          f"\nWrote {o.out} and {csv_path}", file=sys.stderr)


if __name__ == "__main__":
    main()
